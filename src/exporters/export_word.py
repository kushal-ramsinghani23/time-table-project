import os
from docx import Document
from collections import defaultdict

from src.services.services import (
    get_faculty_timetable_for_export,
    get_division_timetable,
    get_lab_timetable,
    get_classroom_occupied
)

# -------------------------------------------------
# CONSTANTS
# -------------------------------------------------
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

TIME_SLOTS = [
    "08:30 - 09:30",
    "09:30 - 10:30",
    "10:30 - 11:30",
    "11:30 - 12:30",
    "12:30 - 13:30",
    "13:30 - 14:30",
    "14:30 - 15:30",
    "15:30 - 16:30",
    "16:30 - 17:30",
    "17:30 - 18:30",
]

BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUTPUT_DIR = os.path.join(BASE_DIR, "data", "output")


# -------------------------------------------------
# HELPERS
# -------------------------------------------------
def ensure_dir(path):
    os.makedirs(path, exist_ok=True)


def merge_entries(entries):
    """
    Same logic as UI:
    Merge multiple timetable entries into one cell
    """
    subjects, divisions, rooms, extras = set(), set(), set(), set()

    for e in entries:
        lines = e.split("\n")
        if len(lines) > 0: subjects.add(lines[0])
        if len(lines) > 1: divisions.add(lines[1])
        if len(lines) > 2: rooms.add(lines[2])
        if len(lines) > 3: extras.add(lines[3])

    result = [
        ", ".join(sorted(subjects)),
        " & ".join(sorted(divisions)),
        ", ".join(sorted(rooms))
    ]

    if extras:
        result.append(", ".join(sorted(extras)))

    return "\n".join(result)


def create_table(doc):
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    hdr[0].text = "Time"
    for i, d in enumerate(DAYS):
        hdr[i + 1].text = d

    return table


def fill_table(table, grid):
    for slot_index, slot in enumerate(TIME_SLOTS, start=1):
        row_cells = table.add_row().cells
        row_cells[0].text = slot

        for i, day in enumerate(DAYS):
            row_cells[i + 1].text = grid.get((day, slot_index), "")


# -------------------------------------------------
# FACULTY TIMETABLE → WORD
# -------------------------------------------------
def export_faculty_timetable(name):
    ensure_dir(os.path.join(OUTPUT_DIR, "faculty_docs"))

    records = get_faculty_timetable_for_export(name)
    cell_map = defaultdict(list)

    for day, slot, subject, room, division, category in records:
        text = f"{subject}\n{division}\n{room}"
        if category == "C":
            text += "\n(Elective)"
        cell_map[(day, slot)].append(text)

    grid = {
        k: merge_entries(v)
        for k, v in cell_map.items()
    }

    doc = Document()
    doc.add_heading("Faculty Timetable", level=1)
    doc.add_paragraph(f"Faculty Name: {name}")

    table = create_table(doc)
    fill_table(table, grid)

    path = os.path.join(OUTPUT_DIR, "faculty_docs", f"{name}_Timetable.docx")
    doc.save(path)
    return path


# -------------------------------------------------
# DIVISION TIMETABLE → WORD
# -------------------------------------------------
def export_division_timetable(division):
    ensure_dir(os.path.join(OUTPUT_DIR, "division_docs"))

    records = get_division_timetable(division)
    cell_map = defaultdict(list)

    for day, slot, _, _, subject, faculty, room, category in records:
        text = f"{subject}\n{faculty}\n{room}"
        if category == "C":
            text += "\n(Elective)"
        cell_map[(day, slot)].append(text)

    grid = {
        k: merge_entries(v)
        for k, v in cell_map.items()
    }

    doc = Document()
    doc.add_heading("Division Timetable", level=1)
    doc.add_paragraph(f"Division: {division}")

    table = create_table(doc)
    fill_table(table, grid)

    path = os.path.join(OUTPUT_DIR, "division_docs", f"{division}_Timetable.docx")
    doc.save(path)
    return path


# -------------------------------------------------
# LAB OCCUPIED → WORD
# -------------------------------------------------
def export_lab_occupied(lab):
    ensure_dir(os.path.join(OUTPUT_DIR, "lab_docs"))

    records = get_lab_timetable(lab)
    cell_map = defaultdict(list)

    for day, slot, _, _, division, batch, subject, faculty in records:
        cell_map[(day, slot)].append(
            f"{subject}\n{division}\n{faculty}"
        )

    grid = {
        k: merge_entries(v)
        for k, v in cell_map.items()
    }

    doc = Document()
    doc.add_heading("Lab Occupancy Timetable", level=1)
    doc.add_paragraph(f"Laboratory: {lab}")

    table = create_table(doc)
    fill_table(table, grid)

    path = os.path.join(OUTPUT_DIR, "lab_docs", f"{lab}_Occupied.docx")
    doc.save(path)
    return path


# -------------------------------------------------
# CLASSROOM OCCUPIED → WORD
# -------------------------------------------------
def export_classroom_occupied(room):
    ensure_dir(os.path.join(OUTPUT_DIR, "classroom_docs"))

    records = get_classroom_occupied(room)
    cell_map = defaultdict(list)

    for day, slot, _, _, division, subject, faculty in records:
        cell_map[(day, slot)].append(
            f"{subject}\n{division}\n{faculty}"
        )

    grid = {
        k: merge_entries(v)
        for k, v in cell_map.items()
    }

    doc = Document()
    doc.add_heading("Classroom Occupancy Timetable", level=1)
    doc.add_paragraph(f"Classroom: {room}")

    table = create_table(doc)
    fill_table(table, grid)

    path = os.path.join(OUTPUT_DIR, "classroom_docs", f"{room}_Occupied.docx")
    doc.save(path)
    return path
